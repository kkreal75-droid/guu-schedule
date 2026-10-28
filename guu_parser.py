"""
Универсальный парсер расписания ГУУ: все файлы с https://guu.ru/student/schedule/ и /student/session/.

Разбор ячеек — по логике parser.py (он остаётся как есть и проверен на МБ(МП) 1-1/1-2):
недели в скобках, «/» между занятиями, объединённая на две строки ячейка = чётные и нечётные.
Поверх неё — то, чего в файле 1 курса не было:
  - шапка (институт / направление / программа / № группы) ищется по подписям в первых колонках
    (A, B или C — по-разному на разных листах), с учётом объединённых ячеек;
  - сетка времени берётся из файла по каждому дню (у ОЗФО, магистратуры и аспирантуры она своя);
  - «раздутые» листы (тысячи пустых строк) читаются только до последней строки сетки;
  - виды «ЛЗ», «(Л 1-11н, ПЗ 13,15н)», «ПЗ (1-12н)», недели без вида «(11 н)», «//» между занятиями;
  - ячейка, объединённая на несколько пар, — занятие в каждой из них;
  - аспирантура: даты и диапазоны дат вместо недель, время и аудитория внутри ячейки;
  - УМБ (магистратура на английском, .doc → .docx через LibreOffice): таблица-календарь по датам;
  - сессия: строки — даты, в ячейке «Предмет (Э) 9-12 Преподаватель Аудитория».

Результат — список групп с парами и отчёт (что не распознано). Идентификаторы групп, имена
и запись файлов — в update.py.
"""

from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from parser import KIND_NAMES, MAX_WEEK, SEMESTER_START, TEACHER, clean, parse_weeks

# ---------- общее ----------

WEEKDAY_NAMES = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
WEEKDAY_SHORT = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

KINDS = {**KIND_NAMES, "ЛЗ": "Лабораторная"}
KIND_TOKEN = r"(?:ЛЗ|ЛР|ПЗ|Л)"

# «/» на отдельной строке (как в parser.py) или «//» между занятиями (магистратура)
PART_SPLIT = re.compile(r"\n\s*/{1,2}\s*\n|\s+//\s*(?:\n|$)|\n\s*//\s+")
# аудитории: «У-432», «ЛК-105/107», «ЦИТ-216», «БЦ-305» + особые места; у языков потока — этаж корпуса
# («ЛК-6 этаж», «У 6 этаж»: аудиторию дают на кафедре)
ROOM = re.compile(r"\b[А-ЯЁ]{1,3}-\d{2,3}(?:/\d{2,3})?|\b[А-ЯЁ]{1,3}[-\s]\d{1,2}\s*этаж|Спортивный комплекс|Актовый зал|ЦК ГУУ")
FLOOR = re.compile(r"^([А-ЯЁ]{1,3})[-\s]+(\d{1,2})\s*этаж$")


def pick_room(rooms: list[str]) -> str | None:
    """Первая аудитория; этаж («ЛК-6 этаж» → «ЛК, 6 этаж») — только если настоящей аудитории в тексте нет
    (в ячейке «(французский) ЛК-6 этаж (китайский) … У-343» остаётся У-343, как раньше)."""
    real = [r for r in rooms if not FLOOR.match(r)]
    if real or not rooms:
        return real[0] if real else None
    m = FLOOR.match(rooms[0])
    return f"{m.group(1)}, {m.group(2)} этаж"
TIME = re.compile(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})")
# в сессии время бывает и без минут: «9-11», «9-10.30», «15-18»
LOOSE_TIME = re.compile(r"(?<![\d.])(\d{1,2})(?:[.:](\d{2}))?\s*[-–—]\s*(\d{1,2})(?:[.:](\d{2}))?(?![\d.])")
DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4})")

# Предметы, у которых в файле обычно нет скобок с неделями, — это не «нераспознанное»
NO_WEEKS_OK = ("проектное обучение", "физическая культура", "элективные дисциплины по физической", "(дпв",
               "иностранный язык", "немецкий", "французский", "китайский", "испанский", "английский")


def norm(s) -> str:
    return clean(s).lower().replace("ё", "е")


def hhmm(h, m) -> str:
    return f"{int(h):02d}:{int(m or 0):02d}"


def time_range(text: str) -> tuple[str, str] | None:
    m = TIME.search(text or "")
    return (hhmm(m[1], m[2]), hhmm(m[3], m[4])) if m else None


def week_of(d: date) -> int:
    return (d - SEMESTER_START).days // 7 + 1


def despace(text: str) -> str:
    """«О Т Р А С Л Е В О Г О    М Е Н Е Д Ж М Е Н Т А» → «ОТРАСЛЕВОГО МЕНЕДЖМЕНТА».
    Слова разделены двумя и более пробелами (или переводом строки), буквы внутри — одним."""
    raw = str(text or "").replace("\xa0", " ").strip()
    tokens = raw.split()
    if len(tokens) >= 3 and sum(len(t) == 1 for t in tokens) >= 0.7 * len(tokens):
        words = [w.replace(" ", "") for w in re.split(r"\s{2,}|\n", raw) if w.strip()]
        return clean(" ".join(words))
    return clean(raw.replace("\n", " "))


def tidy_case(text: str) -> str:
    """«УПРАВЛЕНИЕ ПЕРСОНАЛОМ» → «Управление персоналом» (обычный регистр не трогаем)."""
    letters = [ch for ch in text if ch.isalpha()]
    if letters and sum(ch.isupper() for ch in letters) > 0.8 * len(letters):
        return text[:1].upper() + text[1:].lower()
    return text


INSTITUTES = [
    # короткое имя, полное, по чему узнать в шапке
    ("ИУПСиБК", "Институт управления персоналом, социальных и бизнес-коммуникаций", "УПРАВЛЕНИЯ ПЕРСОНАЛОМ"),
    ("ИИС", "Институт информационных систем", "ИНФОРМАЦИОННЫХ СИСТЕМ"),
    ("ИМ", "Институт маркетинга", "МАРКЕТИНГА"),
    ("ИЭФ", "Институт экономики и финансов", "ЭКОНОМИКИ И ФИНАНСОВ"),
    ("ИГУиП", "Институт государственного управления и права", "ГОСУДАРСТВЕННОГО УПРАВЛЕНИЯ"),
    ("ИОМ", "Институт отраслевого менеджмента", "ОТРАСЛЕВОГО МЕНЕДЖМЕНТА"),
    ("ВШБиТ", "Высшая школа бизнеса и технологий", "ВЫСШАЯ ШКОЛА БИЗНЕСА"),
    # опечатка в шапке сессии ОЗФО: «ГОСУДАРСТВЕННОГО И МУНИЦИПАЛЬНОГО ПРАВА»
    ("ИГУиП", "Институт государственного управления и права", "ГОСУДАРСТВЕННОГО И МУНИЦИПАЛЬНОГО"),
]


def institute_of(text: str) -> tuple[str, str] | None:
    upper = despace(text).upper().replace("Ё", "Е")
    for short, full, mark in INSTITUTES:
        if mark in upper:
            return short, full
    return None


# ---------- результат ----------

@dataclass
class GroupInfo:
    level: str  # bachelor / master / postgrad
    form: str  # ofo / ozfo
    course: int
    institute: str  # короткое: «ИУПСиБК»
    institute_full: str
    direction: str
    program: str
    number: str
    where: str  # «файл · лист · колонка» — для отчёта


@dataclass
class ParsedGroup:
    info: GroupInfo
    lessons: list[dict] = field(default_factory=list)
    grid: dict[int, list[str]] = field(default_factory=dict)  # день недели → ["09:00-10:30", …]


@dataclass
class SessionGroup:
    info: GroupInfo
    title: str  # «Летняя зачётно-экзаменационная сессия 2025/2026»
    academic_year: str  # «2025/2026»
    items: list[dict] = field(default_factory=list)


@dataclass
class Report:
    sheets: list[dict] = field(default_factory=list)  # по листам: группы, пары
    unrecognized: list[dict] = field(default_factory=list)  # ячейки, разобранные «как есть»
    warnings: list[str] = field(default_factory=list)  # где парсер не уверен

    def unknown(self, where: str, text: str, why: str):
        if len(self.unrecognized) < 400:
            self.unrecognized.append({"where": where, "text": clean(text.replace("\n", " "))[:200], "why": why})


# ---------- ячейка с занятием ----------

def _lesson_parens(text: str) -> list[dict]:
    """Скобки с видом и неделями: «(ПЗ 2-16 н)», «(Л 1-11н, ПЗ 13,15н)», «(11 н)», «ПЗ (1-12н)»."""
    out = []
    for m in re.finditer(r"\(([^()]*)\)", text):
        inner = m.group(1).strip()
        # внутри — только виды, числа, «-», «,», «н»: «ПЗ 1-4н, 6-14н», «Л 1-11н, ПЗ 13,15н», «8.14 н»
        if not re.fullmatch(rf"(?:{KIND_TOKEN}|[\d\s,;.\-–нН])+", inner) or not re.search(rf"\d|^{KIND_TOKEN}", inner):
            continue
        if not re.search(r"\d", inner) and not re.fullmatch(rf"{KIND_TOKEN}", inner):
            continue
        # вид прямо перед скобкой: «Математика ПЗ (1-12н)»
        before = re.search(rf"\b({KIND_TOKEN})\s*$", text[: m.start()])
        start = before.start() if before and not re.match(KIND_TOKEN, inner) else m.start()
        chunks = []
        for chunk in re.split(rf"(?={KIND_TOKEN}\b)", inner):
            chunk = chunk.strip(" ,;")
            if not chunk:
                continue
            km = re.match(rf"({KIND_TOKEN})\b", chunk)
            kind = km.group(1) if km else (before.group(1) if before else None)
            spec = re.sub(r"[нН]", "", chunk[km.end():] if km else chunk).replace(".", ",").replace(";", ",").strip(" ,")
            chunks.append((kind, spec))
        out.append({"start": start, "end": m.end(), "chunks": chunks})
    return out


def parse_part(text: str) -> list[dict]:
    """Одно занятие из ячейки → одна или несколько записей. Повторяет parser.parse_part
    (предмет — до первой скобки, преподаватель — первый в тексте, аудитория — после своей скобки)
    и понимает больше видов скобок."""
    text = clean(text.replace("\n", " "))
    if not text:
        return []
    parens = _lesson_parens(text)
    teacher_m = TEACHER.search(text)
    teacher = norm_teacher(teacher_m.group(0)) if teacher_m else None

    if not parens:
        rooms = ROOM.findall(text)
        subject = text
        # «Физическая культура и спорт (Спортивный комплекс)» — место и преподаватель не часть названия
        for found in ([teacher_m.group(0)] if teacher_m else []) + rooms:
            subject = subject.replace(found, " ")
        subject = clean(re.sub(r"\(\s*\)", "", subject)).strip(" ,;/")
        since = re.search(r"\bсо?\s+(\d{1,2})\s*н\b", subject)  # «со 2н»
        weeks = list(range(int(since.group(1)), MAX_WEEK + 1)) if since else None
        if since:
            subject = clean(subject[: since.start()])
        return [dict(subject=subject or text, kind=None, weeks=weeks, explicit=[], teacher=teacher,
                     room=pick_room(rooms), fallback=True)]

    subject = clean(text[: parens[0]["start"]]).strip(" ,;")
    out = []
    for i, p in enumerate(parens):
        tail_end = parens[i + 1]["start"] if i + 1 < len(parens) else len(text)
        rooms = ROOM.findall(text[p["end"]:tail_end])
        if not rooms and len(parens) == 1:
            rooms = ROOM.findall(text)
        for kind, spec in p["chunks"]:
            out.append(dict(
                subject=subject,
                kind=KINDS.get(kind) if kind else None,
                weeks=parse_weeks(spec) or None,
                explicit=[int(x) for x in re.split(r"[,\s]+", spec) if x.isdigit()],
                teacher=teacher,
                room=pick_room(rooms),
                fallback=False,
            ))
    return out


def norm_teacher(name: str) -> str:
    """«Денисова И. В.» → «Денисова И.В.»"""
    return re.sub(r"\.\s+(?=[А-ЯЁ]\.)", ".", clean(name))


def all_teachers(text: str) -> list[str]:
    seen = []
    for m in TEACHER.finditer(clean(text.replace("\n", " "))):
        t = norm_teacher(m.group(0))
        if t not in seen:
            seen.append(t)
    return seen


def split_parts(raw: str) -> list[str]:
    return [p for p in PART_SPLIT.split(raw) if clean(p)]


# ---------- объединённые ячейки ----------

def merged_ranges(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    """(row, col) → (min_row, min_col, max_row, max_col) объединения"""
    m = {}
    for rng in ws.merged_cells.ranges:
        box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = box
    return m


class Sheet:
    """Значения листа с учётом объединений: у объединённой ячейки значение — у верхней левой."""

    def __init__(self, ws):
        self.ws = ws
        self.merged = merged_ranges(ws)
        self.max_row = ws.max_row
        self.max_col = ws.max_column

    def raw(self, r: int, c: int):
        return self.ws.cell(r, c).value

    def value(self, r: int, c: int):
        box = self.merged.get((r, c))
        return self.ws.cell(box[0], box[1]).value if box else self.ws.cell(r, c).value

    def find(self, label: str, rows=range(1, 16), cols=range(1, 7)) -> tuple[int, int] | None:
        for r in rows:
            for c in cols:
                if norm(self.raw(r, c)).startswith(label):
                    return r, c
        return None


# ---------- контекст файла ----------

@dataclass
class FileMeta:
    name: str  # имя файла ГУУ
    text: str  # текст ссылки на сайте
    section: str | None  # раздел страницы: бакалавриат / магистратура / аспирантура

    @property
    def level(self) -> str:
        t = norm(f"{self.section or ''} {self.text} {self.name}")
        if "аспирант" in t:
            return "postgrad"
        if "магистр" in t:
            return "master"
        if "бакалавр" in t:
            return "bachelor"
        return "postgrad" if self.section == "аспирантура" else "bachelor"

    @property
    def form(self) -> str:
        return "ozfo" if "озфо" in norm(f"{self.text} {self.name}") else "ofo"

    @property
    def course(self) -> int | None:
        """Курс из ссылки, если он один («3 курс»); «1-5 курс» — нет."""
        m = re.search(r"(?<![\d-])(\d)\s*[-‑]?\s*(?:курс|к\.)", norm(f"{self.text} {self.name}").replace("-курс", " курс"))
        return int(m.group(1)) if m else None


def sheet_course(title: str) -> int | None:
    m = re.search(r"(\d)\s*(?:к\.|курс)", norm(title))
    return int(m.group(1)) if m else None


# ---------- сетка: бакалавриат, магистратура, ОЗФО ----------

def parse_grid_sheet(ws, meta: FileMeta, report: Report) -> list[ParsedGroup]:
    sh = Sheet(ws)
    where_sheet = f"{meta.name} · {ws.title.strip()}"
    day_cell = sh.find("день")
    if not day_cell:
        return []
    day_row, day_col = day_cell
    time_col, week_col = day_col + 1, day_col + 2
    has_parity = norm(sh.raw(day_row, week_col)).startswith("недел")

    labels = {}
    for key, label in (("institute", "институт"), ("direction", "направление"),
                       ("program", "образовательная программа"), ("number", "№ учебной группы")):
        found = sh.find(label, rows=range(1, day_row), cols=range(1, day_col + 1))
        if found:
            labels[key] = found[0]
    if "program" in labels and "number" not in labels and labels["program"] + 1 < day_row:
        labels["number"] = labels["program"] + 1  # подписи нет, номера — строкой ниже программы
    if "program" not in labels or "number" not in labels:
        report.warnings.append(f"{where_sheet}: нет строк «Образовательная программа» / «№ учебной группы» — лист пропущен")
        return []

    course = sheet_course(ws.title) or meta.course
    if course is None:
        report.warnings.append(f"{where_sheet}: не понял курс — лист пропущен")
        return []
    if meta.course and sheet_course(ws.title) and meta.course != sheet_course(ws.title) and meta.level != "bachelor":
        report.warnings.append(f"{where_sheet}: курс в ссылке {meta.course}, на листе {sheet_course(ws.title)} — беру лист")

    # строки сетки: день / время / неделя. «Раздутые» листы: дальше последней метки не идём.
    rows = []  # (row, weekday, start, end, parity) parity: 1 — нечётная, 0 — чётная, None — без деления
    weekday, slot = None, None
    blank_run = 0
    for r in range(day_row + 1, sh.max_row + 1):
        a, b, c = norm(sh.raw(r, day_col)), clean(sh.raw(r, time_col)), norm(sh.raw(r, week_col)).rstrip(".")
        if not (a or b or c):
            blank_run += 1
            if blank_run > 30:
                break
            continue
        blank_run = 0
        if a in WEEKDAY_NAMES:
            weekday = WEEKDAY_NAMES.index(a)
        t = time_range(b)
        if t:
            slot = t
        if weekday is None or slot is None:
            continue
        if has_parity:
            if c not in ("нечет", "чет"):
                continue
            rows.append((r, weekday, slot[0], slot[1], 1 if c == "нечет" else 0))
        elif t:
            rows.append((r, weekday, slot[0], slot[1], None))

    grid: dict[int, list[str]] = {}
    for _, wd, start, end, _ in rows:
        grid.setdefault(wd, [])
        if f"{start}-{end}" not in grid[wd]:
            grid[wd].append(f"{start}-{end}")
    row_at = {r: (wd, start, end, parity) for r, wd, start, end, parity in rows}

    # колонки групп: там, где в строке «№ учебной группы» стоит номер
    groups: list[ParsedGroup] = []
    last_inst = last_dir = ""
    for col in range(week_col + 1 if has_parity else time_col + 1, sh.max_col + 1):
        inst_raw = despace(sh.value(labels["institute"], col)) if "institute" in labels else ""
        dir_raw = despace(sh.value(labels["direction"], col)) if "direction" in labels else ""
        # у объединённой шапки соседние колонки пустые — это «то же, что слева»
        last_inst = inst_raw or last_inst
        last_dir = dir_raw or last_dir
        num = clean(sh.value(labels["number"], col))
        if isinstance(sh.value(labels["number"], col), float):
            num = str(int(sh.value(labels["number"], col)))
        program = despace(sh.value(labels["program"], col))
        if not re.fullmatch(r"\d{1,2}", num) or not program:
            continue
        inst = institute_of(last_inst) or institute_of(ws.title) or institute_of(program)
        if not inst:
            report.warnings.append(f"{where_sheet}: не понял институт «{last_inst}» (колонка {openpyxl.utils.get_column_letter(col)})")
            inst = (clean(ws.title).split()[0], tidy_case(last_inst) or ws.title)
        info = GroupInfo(level=meta.level, form=meta.form, course=course, institute=inst[0], institute_full=inst[1],
                         direction=tidy_case(last_dir), program=tidy_case(program), number=num,
                         where=f"{where_sheet} · {openpyxl.utils.get_column_letter(col)}")
        group = ParsedGroup(info=info, grid=grid)
        parse_grid_column(sh, col, rows, row_at, group, report)
        groups.append(group)

    report.sheets.append({"file": meta.name, "sheet": ws.title.strip(), "kind": "schedule", "groups": len(groups),
                          "lessons": sum(len(g.lessons) for g in groups)})
    return groups


def parse_grid_column(sh: Sheet, col: int, rows, row_at, group: ParsedGroup, report: Report) -> None:
    col_letter = openpyxl.utils.get_column_letter(col)
    for r, weekday, start, end, parity in rows:
        box = sh.merged.get((r, col))
        if box:
            if r != box[0]:
                continue  # уже обработали на верхней строке объединения
            value = sh.raw(box[0], box[1])
            # какие пары (и чётности) накрывает объединение
            covered: dict[tuple, set] = {}
            for rr in range(box[0], box[2] + 1):
                if rr in row_at:
                    wd, s, e, p = row_at[rr]
                    covered.setdefault((wd, s, e), set()).add(p)
            slots = [(wd, s, e, None if len(ps) >= 2 or None in ps else next(iter(ps))) for (wd, s, e), ps in covered.items()]
        else:
            value = sh.raw(r, col)
            slots = [(weekday, start, end, parity)]
        raw = str(value or "")
        if not clean(raw):
            continue
        # своё время в ячейке: «19.00-20.30\nТехнологии …»
        own = re.match(r"\s*(\d{1,2})[.:](\d{2})\s*[-–]\s*(\d{1,2})[.:](\d{2})\s*\n", raw)
        if own:
            raw = raw[own.end():]
            wd = slots[0][0]
            slots = [(wd, hhmm(own[1], own[2]), hhmm(own[3], own[4]), None)]
            report.warnings.append(f"{group.info.where} {WEEKDAY_SHORT[wd]}: своё время в ячейке {slots[0][1]}–{slots[0][2]}")
        where = f"{group.info.where}{r}"
        for wd, s, e, par in slots:
            add_cell_lessons(group, raw, wd, s, e, par, where, report)
    _ = col_letter


def add_cell_lessons(group: ParsedGroup, raw: str, weekday: int, start: str, end: str, parity: int | None,
                     where: str, report: Report) -> None:
    for part in split_parts(raw):
        teachers = all_teachers(part)
        for item in parse_part(part):
            if item.pop("fallback") and not any(norm(item["subject"]).startswith(ok) for ok in NO_WEEKS_OK):
                report.unknown(where, part, "нет скобок с видом и неделями — взято как есть, все недели")
            weeks = item["weeks"] or list(range(1, MAX_WEEK + 1))
            if parity is not None:
                # как в parser.py: «3-13» в нечётной строке = 3, 5, …, 13; предупреждаем, только если
                # отброшена явно названная неделя
                explicit = set(item.get("explicit") or [])
                dropped = [w for w in weeks if w % 2 != parity and w in explicit]
                weeks = [w for w in weeks if w % 2 == parity]
                if dropped:
                    report.warnings.append(
                        f"{where} {WEEKDAY_SHORT[weekday]} {start} «{item['subject']}»: недели {dropped} "
                        f"не совпадают с чётностью строки — отброшены")
            if not weeks:
                continue
            item.pop("explicit", None)
            lesson = {"weekday": weekday, "start": start, "end": end, **item, "weeks": weeks}
            if len(teachers) > 1:
                lesson["teachers"] = teachers  # для teachers.json: в ячейке несколько преподавателей
            group.lessons.append(lesson)


# ---------- аспирантура: даты вместо недель ----------

def split_grid(start: str, end: str, grid: list[str]) -> list[tuple[str, str]]:
    """«10:40-14:25» → пары сетки внутри: 10:40–12:10, 12:55–14:25. Нет таких — как есть."""
    inside = [tuple(g.split("-")) for g in grid if g.split("-")[0] >= start and g.split("-")[1] <= end]
    return inside or [(start, end)]


def dates_in(spec: str, weekday: int | None) -> tuple[list[date], list[str]]:
    """«13.10.2026-10.11.2026» (по этому дню недели), «16.11.2026», «12.10, 19.10.2026»."""
    found, notes = [], []
    text = spec.replace("–", "-").replace("—", "-")
    for m in re.finditer(r"(\d{1,2}\.\d{1,2}\.\d{4})\s*-\s*(\d{1,2}\.\d{1,2}\.\d{4})|(\d{1,2}\.\d{1,2}\.\d{4})", text):
        try:
            if m.group(1):
                a = datetime.strptime(m.group(1), "%d.%m.%Y").date()
                b = datetime.strptime(m.group(2), "%d.%m.%Y").date()
                wd = weekday if weekday is not None else a.weekday()
                d = a
                while d <= b:
                    if d.weekday() == wd:
                        found.append(d)
                    d += timedelta(days=1)
            else:
                d = datetime.strptime(m.group(3), "%d.%m.%Y").date()
                if weekday is not None and d.weekday() != weekday:
                    notes.append(f"{d:%d.%m.%Y} — не {WEEKDAY_NAMES[weekday]}, беру день по дате")
                found.append(d)
        except ValueError:
            notes.append(f"не разобрал дату «{m.group(0)}»")
    return found, notes


def parse_postgrad_sheet(ws, meta: FileMeta, report: Report) -> list[ParsedGroup]:
    sh = Sheet(ws)
    where_sheet = f"{meta.name} · {ws.title.strip()}"
    spec_cell = sh.find("научная специальность", rows=range(1, 20), cols=range(1, 4))
    if not spec_cell:
        return []
    spec_row = spec_cell[0]
    sub_row = spec_row + 1  # программы внутри специальности (5.2.3 → несколько)
    course = meta.course or sheet_course(ws.title)
    if meta.course and sheet_course(ws.title) and meta.course != sheet_course(ws.title):
        report.warnings.append(f"{where_sheet}: курс в ссылке {meta.course}, в названии листа «{ws.title.strip()}» — беру ссылку")
    if course is None:
        report.warnings.append(f"{where_sheet}: не понял курс — лист пропущен")
        return []

    # строки: день в A, время в B, по строке на пару
    rows, weekday = [], None
    for r in range(spec_row + 1, sh.max_row + 1):
        a = norm(sh.raw(r, 1))
        if a in WEEKDAY_NAMES:
            weekday = WEEKDAY_NAMES.index(a)
        t = time_range(clean(sh.raw(r, 2)))
        if t and weekday is not None:
            rows.append((r, weekday, t[0], t[1]))
    grid: dict[int, list[str]] = {}
    for _, wd, s, e in rows:
        grid.setdefault(wd, []).append(f"{s}-{e}")
    row_at = {r: (wd, s, e) for r, wd, s, e in rows}

    groups = []
    for col in range(3, sh.max_col + 1):
        head = clean(str(sh.value(spec_row, col) or "").replace("\n", " "))
        sub = clean(str(sh.raw(sub_row, col) or "").replace("\n", " "))
        if not head or norm(head).startswith("научная"):
            continue
        m = re.match(r"(\d+(?:\.\d+){1,2})\.?\s*(.*)", head)
        if not m:
            continue
        code, title = m.group(1), m.group(2).strip()
        info = GroupInfo(level="postgrad", form="ofo", course=course, institute="Аспирантура",
                         institute_full="Аспирантура", direction=f"{code} {title}",
                         program=sub or f"{code} {title}", number="1",
                         where=f"{where_sheet} · {openpyxl.utils.get_column_letter(col)}")
        group = ParsedGroup(info=info, grid=grid)
        for r, wd, s, e in rows:
            box = sh.merged.get((r, col))
            if box and r != box[0]:
                continue
            raw = str((sh.raw(box[0], box[1]) if box else sh.raw(r, col)) or "")
            if not clean(raw):
                continue
            covered = [row_at[rr][1:] for rr in range(box[0], box[2] + 1) if rr in row_at] if box else [(s, e)]
            parse_postgrad_cell(group, raw, wd, covered, f"{info.where}{r}", report)
        groups.append(group)
    report.sheets.append({"file": meta.name, "sheet": ws.title.strip(), "kind": "schedule", "groups": len(groups),
                          "lessons": sum(len(g.lessons) for g in groups)})
    return groups


def parse_postgrad_cell(group: ParsedGroup, raw: str, weekday: int, covered: list[tuple[str, str]],
                        where: str, report: Report) -> None:
    paragraphs = [p for p in re.split(r"\n\s*\n", raw.strip()) if p.strip()]
    subject = clean(paragraphs[0].replace("\n", " "))
    rest = "\n".join(paragraphs[1:])
    teacher_m = TEACHER.search(clean(rest.replace("\n", " ")))
    teacher = norm_teacher(teacher_m.group(0)) if teacher_m else None
    # блоки: строка с датами, за ней — строки «время аудитория»
    blocks: list[dict] = []
    for line in rest.split("\n"):
        line = clean(line)
        if not line:
            continue
        if DATE.search(line):
            blocks.append({"dates": line, "times": [], "room": None})
            rm = ROOM.findall(DATE.sub("", line))
            if rm:
                blocks[-1]["room"] = rm[0]
            continue
        if not blocks:
            continue
        t = time_range(line)
        rooms = ROOM.findall(line)
        if t:
            blocks[-1]["times"].append((t[0], t[1], rooms[0] if rooms else None))
        elif rooms:
            blocks[-1]["room"] = rooms[0]
    if not blocks:
        report.unknown(where, raw, "аспирантура: нет дат — пропущено")
        return
    grid_day = group.grid.get(weekday, [])
    acc: dict[tuple, set] = {}
    for block in blocks:
        days, notes = dates_in(block["dates"], weekday)
        for note in notes:
            report.warnings.append(f"{where}: {note}")
        slots = []
        if block["times"]:
            for s, e, room in block["times"]:
                slots += [(ss, ee, room or block["room"]) for ss, ee in split_grid(s, e, grid_day)]
        else:
            slots = [(s, e, block["room"]) for s, e in covered]
        for d in days:
            w = week_of(d)
            if w < 1:
                continue
            for s, e, room in slots:
                acc.setdefault((d.weekday(), s, e, room), set()).add(w)
    for (wd, s, e, room), weeks in acc.items():
        group.lessons.append({"weekday": wd, "start": s, "end": e, "subject": subject, "kind": None,
                              "weeks": sorted(weeks), "teacher": teacher, "room": room})


# ---------- .docx (УМБ: и расписание, и сессия — таблица-календарь) ----------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _text(el) -> str:
    lines = []
    for p in el.iter(W + "p"):
        lines.append("".join((t.text or "") for t in p.iter(W + "t")))
    return "\n".join(lines).strip()


def read_docx(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    """Абзацы вне таблиц и таблицы (строки → ячейки по сетке, с учётом gridSpan и vMerge)."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(W + "body")
    paragraphs, tables = [], []
    for el in body:
        if el.tag == W + "p":
            t = _text(el)
            if t:
                paragraphs.append(t)
        elif el.tag == W + "tbl":
            grid_rows: list[list[str]] = []
            for tr in el.findall(W + "tr"):
                row: list[str] = []
                for tc in tr.findall(W + "tc"):
                    pr = tc.find(W + "tcPr")
                    span = 1
                    vmerge_continue = False
                    if pr is not None:
                        gs = pr.find(W + "gridSpan")
                        if gs is not None:
                            span = int(gs.get(W + "val", "1"))
                        vm = pr.find(W + "vMerge")
                        if vm is not None and vm.get(W + "val") in (None, "continue"):
                            vmerge_continue = True
                    text = _text(tc)
                    if vmerge_continue and grid_rows and len(grid_rows[-1]) > len(row):
                        text = grid_rows[-1][len(row)]
                    row.extend([text] * span)
                grid_rows.append(row)
            tables.append(grid_rows)
            # абзацы внутри первой ячейки-шапки тоже нужны (институт, программа)
            for r in grid_rows[:1]:
                if len(set(r)) == 1 and r and r[0]:
                    paragraphs.extend(r[0].split("\n"))
    return paragraphs, tables


def docx_group_info(paragraphs: list[str], meta: FileMeta, report: Report) -> GroupInfo | None:
    text = "\n".join(paragraphs)
    inst = institute_of(text)
    course_m = re.search(r"(\d)\s*курса", text)
    prog_m = re.search(r"образовательная программа\s*[«\"]([^»\"]+)[»\"]\s*(\([^)]*\))?", text, re.I)
    dir_m = re.search(r"направлени[ея] подготовки\s*[«\"]([^»\"]+)[»\"]", text, re.I)
    if not (inst and prog_m):
        report.warnings.append(f"{meta.name}: не нашёл институт или программу в шапке документа — файл пропущен")
        return None
    program = clean(prog_m.group(1) + (" " + prog_m.group(2) if prog_m.group(2) else ""))
    course = int(course_m.group(1)) if course_m else meta.course
    if course is None:
        report.warnings.append(f"{meta.name}: не понял курс — файл пропущен")
        return None
    level = "master" if re.search(r"магистратур", text, re.I) else meta.level
    return GroupInfo(level=level, form="ozfo" if re.search(r"очно-заочн", text, re.I) else "ofo", course=course,
                     institute=inst[0], institute_full=inst[1], direction=clean(dir_m.group(1)) if dir_m else "",
                     program=program, number="1", where=meta.name)


def docx_room(paragraphs: list[str]) -> str | None:
    text = " ".join(paragraphs)
    m = re.search(r"БЦ-?\s*(\d{2,4})", text) or re.search(r"Бизнес[\s-]*центр[^\d]{0,20}(\d{2,4})", text, re.I)
    return f"БЦ-{m.group(1)}" if m else None


HEAD_DAY = re.compile(r"(понедельник|вторник|среда|четверг|пятница|суббота)\s*(\d{1,2}\.\d{1,2}\.\d{4})?", re.I)

# стандартная сетка ГУУ: «18:50-22:00» = две вечерние пары
STANDARD_GRID = ["09:00-10:30", "10:40-12:10", "12:55-14:25", "14:35-16:05", "16:15-17:45", "17:55-19:25",
                 "18:50-20:20", "20:30-22:00"]


def _fix_row_dates(cells: list[tuple[int, int, date | None]], where: str, report: Report) -> dict[int, date]:
    """Даты в строке-заголовке недели: (колонка, день недели, дата). Опечатки («31.10.2025» в
    субботу 2026 года, «07.11» вместо «07.12») исправляем по остальным дням строки."""
    mondays = [d - timedelta(days=wd) for _, wd, d in cells if d and d.weekday() == wd]
    monday = max(set(mondays), key=mondays.count) if mondays else None
    out = {}
    for col, wd, d in cells:
        good = monday + timedelta(days=wd) if monday else d
        if d and good and d != good:
            report.warnings.append(f"{where}: дата {d:%d.%m.%Y} ({WEEKDAY_NAMES[wd]}) похожа на опечатку — беру {good:%d.%m.%Y}")
        if good:
            out[col] = good
    return out


def parse_calendar_docx(path: Path, meta: FileMeta, report: Report) -> list[ParsedGroup]:
    """УМБ: строки «понедельник 05.10.2026 | вторник 06.10.2026 | …», под ними — пары на эти дни."""
    paragraphs, tables = read_docx(path)
    info = docx_group_info(paragraphs, meta, report)
    if not info:
        return []
    room = docx_room(paragraphs)
    group = ParsedGroup(info=info)
    acc: dict[tuple, set] = {}
    for table in tables:
        dates: dict[int, date] = {}
        for ri, row in enumerate(table):
            heads = [(ci, HEAD_DAY.search(cell)) for ci, cell in enumerate(row)]
            heads = [(ci, m) for ci, m in heads if m and m.group(2)]
            if len(heads) >= 2:
                parsed = []
                for ci, m in heads:
                    try:
                        d = datetime.strptime(m.group(2), "%d.%m.%Y").date()
                    except ValueError:
                        d = None
                    parsed.append((ci, WEEKDAY_NAMES.index(m.group(1).lower()), d))
                dates = _fix_row_dates(parsed, f"{meta.name} · строка {ri + 1}", report)
                continue
            for ci, cell in enumerate(row):
                if ci not in dates or not clean(cell) or (ci > 0 and row[ci - 1] == cell and ci - 1 in dates):
                    continue
                parse_calendar_cell(acc, cell, dates[ci], room, f"{meta.name} · строка {ri + 1}", report)
    for (wd, s, e, subject, teacher, rm), weeks in acc.items():
        group.lessons.append({"weekday": wd, "start": s, "end": e, "subject": subject, "kind": None,
                              "weeks": sorted(weeks), "teacher": teacher, "room": rm})
    group.grid = {wd: sorted({f"{l['start']}-{l['end']}" for l in group.lessons if l["weekday"] == wd})
                  for wd in range(7)}
    group.grid = {wd: g for wd, g in group.grid.items() if g}
    report.sheets.append({"file": meta.name, "sheet": "таблица", "kind": "schedule", "groups": 1,
                          "lessons": len(group.lessons)})
    return [group]


def parse_calendar_cell(acc, cell: str, day: date, room: str | None, where: str, report: Report) -> None:
    times, text_lines = [], []
    for line in cell.split("\n"):
        line = clean(line)
        if not line:
            continue
        found = list(TIME.finditer(line))
        if found:
            times += [(hhmm(m[1], m[2]), hhmm(m[3], m[4])) for m in found]
            rest = clean(TIME.sub("", line))
            if rest:
                text_lines.append(rest)
        else:
            text_lines.append(line)
    if not text_lines:
        return
    joined = " ".join(text_lines)
    teacher_m = TEACHER.search(joined)
    teacher = norm_teacher(teacher_m.group(0)) if teacher_m else None
    subject = text_lines[0]
    if teacher_m and subject.startswith(teacher_m.group(0)):
        subject = joined
    if not times:
        report.unknown(where, cell, "нет времени — пропущено")
        return
    w = week_of(day)
    for s, e in times:
        for ss, ee in split_grid(s, e, STANDARD_GRID):
            acc.setdefault((day.weekday(), ss, ee, subject, teacher, room), set()).add(w)


# ---------- сессия ----------

SESSION_KINDS = {"Э": "Экзамен", "З": "Зачёт", "ЗО": "Зачёт с оценкой", "ДЗ": "Зачёт с оценкой",
                 "КР": "Курсовая работа", "КП": "Курсовой проект", "К": "Консультация"}


def academic_year(d: date) -> str:
    y = d.year if d.month >= 8 else d.year - 1
    return f"{y}/{y + 1}"


def session_time(text: str) -> tuple[str, str] | None:
    m = LOOSE_TIME.search(text)
    if not m:
        return None
    a, b = hhmm(m[1], m[2]), hhmm(m[3], m[4])
    return (a, b) if "07:00" <= a < b <= "23:00" else None


TITLES = re.compile(r"\b(?:[кд]\.\s?[а-я]{1,4}\.\s?н\.|доцент|профессор|ст\.\s?препод\.?|старший преподаватель|преподаватель|ассистент)[,.]?", re.I)


def parse_session_text(text: str, day: date, until: date | None = None) -> list[dict]:
    items = []
    for part in split_parts(text):
        flat = despace(part)  # «У Ч Е Б Н А Я   П Р А К Т И К А» → «УЧЕБНАЯ ПРАКТИКА»
        if not flat:
            continue
        t = session_time(flat)
        rest = clean(LOOSE_TIME.sub(" ", flat, count=1)) if t else flat
        rest = clean(TITLES.sub(" ", rest))
        kind_m = re.search(r"\((ЗО|ДЗ|КР|КП|Э|З|К)\)", rest)
        teacher_m = TEACHER.search(rest)
        rooms = ROOM.findall(flat)
        cut = min([x for x in (kind_m.start() if kind_m else None,
                               teacher_m.start() if teacher_m else None,
                               rest.find(rooms[0]) if rooms and rooms[0] in rest else None) if x is not None] or [len(rest)])
        subject = clean(rest[:cut]).strip(" ,;")
        subject = subject[:1].upper() + subject[1:].lower() if subject.isupper() and len(subject) > 3 else subject
        if not subject:
            continue
        if re.search(r"консультац", subject, re.I) and not kind_m:
            kind = "Консультация"
        else:
            kind = SESSION_KINDS.get(kind_m.group(1)) if kind_m else None
        items.append({"date": day.isoformat(), **({"until": until.isoformat()} if until and until > day else {}),
                      "start": t[0] if t else None, "end": t[1] if t else None,
                      "subject": subject, "kind": kind, "teacher": norm_teacher(teacher_m.group(0)) if teacher_m else None,
                      "room": " / ".join(dict.fromkeys(rooms)) if rooms else None})
    return items


def parse_session_sheet(ws, meta: FileMeta, report: Report) -> list[SessionGroup]:
    sh = Sheet(ws)
    where_sheet = f"{meta.name} · {ws.title.strip()}"
    date_cell = sh.find("дата", rows=range(1, 16), cols=range(1, 6))
    if not date_cell:
        return []
    date_row, date_col = date_cell
    labels = {}
    for key, label in (("institute", "институт"), ("direction", "направление"),
                       ("program", "образовательная программа"), ("number", "№ учебной группы")):
        found = sh.find(label, rows=range(1, date_row + 1), cols=range(1, date_col + 2))
        if found:
            labels[key] = found[0]
    if "program" in labels and "number" not in labels and labels["program"] + 1 < date_row:
        labels["number"] = labels["program"] + 1  # подписи нет, номера — строкой ниже программы
    if "program" not in labels or "number" not in labels:
        report.warnings.append(f"{where_sheet}: нет шапки групп — лист сессии пропущен")
        return []
    title = clean(" ".join(str(sh.raw(r, c) or "") for r in range(1, 4) for c in range(1, sh.max_col + 1)))
    title_m = re.search(r"расписание\s+(.+?сессии)\s+(\d{4}/\d{4})", title, re.I)
    course = sheet_course(ws.title) or meta.course
    if course is None:
        report.warnings.append(f"{where_sheet}: не понял курс — лист сессии пропущен")
        return []

    # строки с датой
    day_rows = []
    blank_run = 0
    for r in range(date_row + 1, sh.max_row + 1):
        v = sh.raw(r, date_col)
        if isinstance(v, datetime):
            day_rows.append((r, v.date()))
            blank_run = 0
        elif isinstance(v, str) and DATE.search(v):
            m = DATE.search(v)
            day_rows.append((r, date(int(m[3]), int(m[2]), int(m[1]))))
            blank_run = 0
        else:
            blank_run += 1
            if blank_run > 20:
                break
    if not day_rows:
        return []
    year = academic_year(day_rows[0][1])
    season = "Зимняя" if day_rows[0][1].month in (12, 1, 2) else "Летняя"
    session_title = f"{season} сессия {year}" if not title_m else f"{tidy_case(title_m.group(1))} {title_m.group(2)}"

    groups = []
    last_inst = last_dir = ""
    for col in range(date_col + 1, sh.max_col + 1):
        inst_raw = despace(sh.value(labels["institute"], col)) if "institute" in labels else ""
        dir_raw = despace(sh.value(labels["direction"], col)) if "direction" in labels else ""
        last_inst = inst_raw or last_inst
        last_dir = dir_raw or last_dir
        num_v = sh.value(labels["number"], col)
        num = str(int(num_v)) if isinstance(num_v, float) else clean(num_v)
        program = despace(sh.value(labels["program"], col))
        if not re.fullmatch(r"\d{1,2}", num) or not program:
            continue
        inst = institute_of(last_inst) or institute_of(ws.title)
        if not inst:
            report.warnings.append(f"{where_sheet}: не понял институт «{last_inst}»")
            continue
        info = GroupInfo(level=meta.level, form=meta.form, course=course, institute=inst[0], institute_full=inst[1],
                         direction=tidy_case(last_dir), program=tidy_case(program), number=num,
                         where=f"{where_sheet} · {openpyxl.utils.get_column_letter(col)}")
        sg = SessionGroup(info=info, title=session_title, academic_year=year)
        day_at = dict(day_rows)
        for r, d in day_rows:
            box = sh.merged.get((r, col))
            if box and (box[2] > box[0]) and r != box[0]:
                continue
            raw = str((sh.raw(box[0], box[1]) if box else sh.raw(r, col)) or "")
            until = max((day_at[rr] for rr in range(box[0], box[2] + 1) if rr in day_at), default=None) if box else None
            if clean(raw):
                sg.items += parse_session_text(raw, d, until)
        groups.append(sg)
    report.sheets.append({"file": meta.name, "sheet": ws.title.strip(), "kind": "session", "groups": len(groups),
                          "items": sum(len(g.items) for g in groups), "title": session_title})
    return groups


def parse_session_docx(path: Path, meta: FileMeta, report: Report) -> list[SessionGroup]:
    paragraphs, tables = read_docx(path)
    info = docx_group_info(paragraphs, meta, report)
    if not info:
        return []
    room = docx_room(paragraphs)
    items = []
    for table in tables:
        for ri, row in enumerate(table):
            seen = set()
            for cell in row:
                if cell in seen:
                    continue
                seen.add(cell)
                m = DATE.search(cell)
                if not m:
                    continue
                rest = clean(DATE.sub("", cell, count=1))
                if not rest:
                    continue
                day = date(int(m[3]), int(m[2]), int(m[1]))
                for item in parse_session_text(cell[m.end():], day):
                    item["room"] = item["room"] or room
                    items.append(item)
    if not items:
        return []
    first = min(date.fromisoformat(i["date"]) for i in items)
    year = academic_year(first)
    season = "Зимняя" if first.month in (12, 1, 2) else "Летняя"
    sg = SessionGroup(info=info, title=f"{season} сессия {year}", academic_year=year, items=items)
    report.sheets.append({"file": meta.name, "sheet": "таблица", "kind": "session", "groups": 1, "items": len(items),
                          "title": sg.title})
    return [sg]


# ---------- файл целиком ----------

def parse_schedule_file(path: Path, meta: FileMeta, report: Report) -> list[ParsedGroup]:
    if path.suffix.lower() == ".docx":
        return parse_calendar_docx(path, meta, report)
    wb = openpyxl.load_workbook(path)
    groups: list[ParsedGroup] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        if Sheet(ws).find("научная специальность", rows=range(1, 20), cols=range(1, 4)):
            groups += parse_postgrad_sheet(ws, meta, report)
        else:
            groups += parse_grid_sheet(ws, meta, report)
    return groups


def parse_session_file(path: Path, meta: FileMeta, report: Report) -> list[SessionGroup]:
    if path.suffix.lower() == ".docx":
        return parse_session_docx(path, meta, report)
    wb = openpyxl.load_workbook(path)
    out: list[SessionGroup] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        if Sheet(ws).find("научная специальность", rows=range(1, 20), cols=range(1, 4)):
            out += parse_postgrad_session_sheet(ws, meta, report)
        else:
            out += parse_session_sheet(ws, meta, report)
    return out


def parse_postgrad_session_sheet(ws, meta: FileMeta, report: Report) -> list[SessionGroup]:
    """Сессия аспирантуры: та же сетка по дням недели, в ячейке «Предмет / дата время / ФИО / ауд.»."""
    sh = Sheet(ws)
    spec_row = sh.find("научная специальность", rows=range(1, 20), cols=range(1, 4))[0]
    course = meta.course or sheet_course(ws.title)
    if course is None:
        report.warnings.append(f"{meta.name} · {ws.title.strip()}: не понял курс — лист сессии пропущен")
        return []
    groups = []
    for col in range(2, sh.max_col + 1):
        head = clean(str(sh.value(spec_row, col) or "").replace("\n", " "))
        sub = clean(str(sh.raw(spec_row + 1, col) or "").replace("\n", " "))
        m = re.match(r"(\d+(?:\.\d+){1,2})\.?\s*(.*)", head)
        if not m:
            continue
        code, title = m.group(1), m.group(2).strip()
        info = GroupInfo(level="postgrad", form="ofo", course=course, institute="Аспирантура",
                         institute_full="Аспирантура", direction=f"{code} {title}",
                         program=sub or f"{code} {title}", number="1",
                         where=f"{meta.name} · {ws.title.strip()} · {openpyxl.utils.get_column_letter(col)}")
        items = []
        seen = set()
        for r in range(spec_row + 2, sh.max_row + 1):
            box = sh.merged.get((r, col))
            key = box[:2] if box else (r, col)
            if key in seen:
                continue
            seen.add(key)
            raw = str((sh.raw(*key)) or "")
            if not clean(raw):
                continue
            paragraphs = [p for p in re.split(r"\n\s*\n", raw.strip()) if p.strip()]
            subject = clean(paragraphs[0].replace("\n", " "))
            rest = "\n".join(paragraphs[1:])
            flat = clean(rest.replace("\n", " "))
            teacher_m = TEACHER.search(flat)
            rooms = ROOM.findall(flat)
            for line in rest.split("\n"):
                dm = DATE.search(line)
                if not dm:
                    continue
                t = session_time(DATE.sub(" ", line))
                items.append({"date": date(int(dm[3]), int(dm[2]), int(dm[1])).isoformat(),
                              "start": t[0] if t else None, "end": t[1] if t else None, "subject": subject,
                              "kind": None, "teacher": norm_teacher(teacher_m.group(0)) if teacher_m else None,
                              "room": rooms[0] if rooms else None})
        if not items:
            continue
        first = min(date.fromisoformat(i["date"]) for i in items)
        season = "Зимняя" if first.month in (12, 1, 2) else "Летняя"
        groups.append(SessionGroup(info=info, title=f"{season} сессия {academic_year(first)}",
                                   academic_year=academic_year(first), items=sorted(items, key=lambda i: i["date"])))
    report.sheets.append({"file": meta.name, "sheet": ws.title.strip(), "kind": "session", "groups": len(groups),
                          "items": sum(len(g.items) for g in groups)})
    return groups
